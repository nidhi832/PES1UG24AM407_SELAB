import json
import re
import os
import html
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
    KeepTogether,
)

def clean_dialogue(transcript_path):
    with open(transcript_path, 'r', encoding='utf-8') as f:
        items = [json.loads(line) for line in f]

    dialogue = []
    for item in items:
        t = item.get('type')
        if t == 'USER_INPUT':
            raw = item.get('content', '')
            m = re.search(r'<USER_REQUEST>\s*(.*?)\s*</USER_REQUEST>', raw, re.DOTALL)
            content = m.group(1).strip() if m else raw.strip()
            # Remove any trailing metadata tags
            content = re.sub(r'<ADDITIONAL_METADATA>.*?</ADDITIONAL_METADATA>', '', content, flags=re.DOTALL).strip()
            if content:
                dialogue.append(('User', content, item.get('created_at', '')))
        elif t == 'PLANNER_RESPONSE':
            content = item.get('content', '')
            if content and content.strip() and not content.startswith('Wait for background task'):
                dialogue.append(('AI Assistant (Antigravity)', content.strip(), item.get('created_at', '')))
    return dialogue

def generate_docx(dialogue, output_path):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run('Lab 4: VibeCoding — Pair Programming Chat History')
    title_run.font.name = 'Segoe UI'
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 60, 120)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run('Lab Assignment: 44_simple-platformer | Repository: nidhi832/44_simple-platformer_SE_LAB4\nPair Programmer Model: Google Antigravity (Gemini 3.7 Flash)')
    sub_run.font.name = 'Segoe UI'
    sub_run.font.size = Pt(10.5)
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_borders(cell, border_color):
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(borders)

    def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    for i, (role, text, time_str) in enumerate(dialogue):
        is_user = (role == 'User')
        bg_color = 'EEF4FF' if is_user else 'F8FAFC'
        border_color = '2563EB' if is_user else '059669'
        role_color = RGBColor(37, 99, 235) if is_user else RGBColor(5, 150, 105)

        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.8)

        cell = table.cell(0, 0)
        set_cell_background(cell, bg_color)
        set_cell_borders(cell, border_color)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        header_p = cell.paragraphs[0]
        header_p.paragraph_format.space_before = Pt(2)
        header_p.paragraph_format.space_after = Pt(4)
        r_role = header_p.add_run(f'Message #{i+1} • {role}')
        r_role.font.name = 'Segoe UI'
        r_role.font.bold = True
        r_role.font.size = Pt(11)
        r_role.font.color.rgb = role_color

        if time_str:
            clean_time = time_str[:19].replace('T', ' ')
            r_time = header_p.add_run(f'   [{clean_time}]')
            r_time.font.name = 'Segoe UI'
            r_time.font.size = Pt(9)
            r_time.font.color.rgb = RGBColor(140, 140, 140)

        lines = text.split('\n')
        in_code = False
        for line in lines:
            trimmed = line.strip()
            if trimmed.startswith('```'):
                in_code = not in_code
                continue
            
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            run = p.add_run(line)
            if in_code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
            else:
                run.font.name = 'Segoe UI'
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(30, 30, 30)

        sep = doc.add_paragraph()
        sep.paragraph_format.space_after = Pt(6)

    doc.save(output_path)
    print(f'DOCX saved successfully: {output_path}')

def generate_pdf(dialogue, output_path):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E3A8A'),
        spaceAfter=4,
    )

    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#6B7280'),
        spaceAfter=15,
    )

    user_header_style = ParagraphStyle(
        'UserHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#1D4ED8'),
    )

    ai_header_style = ParagraphStyle(
        'AIHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#047857'),
    )

    body_style = ParagraphStyle(
        'MsgBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#1F2937'),
    )

    code_style = ParagraphStyle(
        'CodeBody',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor('#0F172A'),
    )

    story = []
    story.append(Paragraph('Lab 4: VibeCoding — Pair Programming Chat History', title_style))
    story.append(Paragraph('Repository: nidhi832/44_simple-platformer_SE_LAB4 | Pair Programmer: Google Antigravity (Gemini 3.7 Flash)', sub_style))
    story.append(HRFlowable(width='100%', thickness=1, color=colors.HexColor('#E5E7EB'), spaceAfter=12))

    for i, (role, text, time_str) in enumerate(dialogue):
        is_user = (role == 'User')
        bg_col = colors.HexColor('#EFF6FF') if is_user else colors.HexColor('#F9FAFB')
        border_col = colors.HexColor('#3B82F6') if is_user else colors.HexColor('#10B981')
        h_style = user_header_style if is_user else ai_header_style

        time_display = f" <font color='#9CA3AF' size='7.5'>[{time_str[:19].replace('T', ' ')}]</font>" if time_str else ""
        header_p = Paragraph(f"<b>Message #{i+1} • {role}</b>{time_display}", h_style)

        story.append(HRFlowable(width='100%', thickness=1.5, color=border_col, spaceBefore=12, spaceAfter=5))
        story.append(header_p)
        story.append(Spacer(1, 4))

        # Process lines safely escaping HTML entities
        lines = text.split('\n')
        in_code = False
        code_acc = []
        for line in lines:
            trimmed = line.strip()
            if trimmed.startswith('```'):
                if in_code and code_acc:
                    code_text = html.escape('\n'.join(code_acc)).replace(' ', '&nbsp;').replace('\n', '<br/>')
                    story.append(Paragraph(code_text, code_style))
                    code_acc = []
                in_code = not in_code
                continue

            if in_code:
                code_acc.append(line)
            else:
                if trimmed:
                    escaped_line = html.escape(line)
                    # convert markdown bold **text** to <b>text</b>
                    escaped_line = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', escaped_line)
                    # convert `code` to font
                    escaped_line = re.sub(r'`(.*?)`', r'<font face="Courier">\1</font>', escaped_line)
                    story.append(Paragraph(escaped_line, body_style))
                else:
                    story.append(Spacer(1, 3))

        if code_acc:
            code_text = html.escape('\n'.join(code_acc)).replace(' ', '&nbsp;').replace('\n', '<br/>')
            story.append(Paragraph(code_text, code_style))

        story.append(Spacer(1, 6))

    doc.build(story)
    print(f'PDF saved successfully: {output_path}')

if __name__ == '__main__':
    t_path = r'C:\Users\admin\.gemini\antigravity\brain\1b198803-f9ea-4505-a1b1-0cb1b1e469e0\.system_generated\logs\transcript_full.jsonl'
    out_dir = r'C:\Users\admin\Desktop\SE LAB 4\Lab-4'
    os.makedirs(out_dir, exist_ok=True)

    dialogue = clean_dialogue(t_path)
    print(f'Loaded {len(dialogue)} dialogue items.')

    docx_path = os.path.join(out_dir, 'Lab4_Chat_History.docx')
    pdf_path = os.path.join(out_dir, 'Lab4_Chat_History.pdf')

    generate_docx(dialogue, docx_path)
    generate_pdf(dialogue, pdf_path)
